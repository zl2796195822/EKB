"""Real OOXML parser coverage for the PH3 parser registry."""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from ekb_api.services.parsers.registry import (
    FILE_EMPTY,
    MIME_UNSUPPORTED,
    PARSER_CORRUPT,
    ParserError,
    ParserRegistry,
)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "销售数据"
    sheet.append(["客户", "金额"])
    sheet.append(["甲公司", 1200])
    second_sheet = workbook.create_sheet("说明")
    second_sheet["A1"] = "本表由业务团队维护"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "季度报告"
    slide.placeholders[1].text = "收入保持增长"
    slide.notes_slide.notes_text_frame.text = "演讲者备注"

    table_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = table_slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(5), Inches(2)).table
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "订单"
    table.cell(1, 1).text = "42"

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def test_registry_resolves_and_parses_real_xlsx() -> None:
    result = ParserRegistry().parse(_xlsx_bytes(), filename="sales.xlsx", mime=XLSX_MIME)

    assert result.metadata == {
        "format": "xlsx",
        "sheets": 2,
        "populated_sheets": 2,
        "sections": 3,
    }
    assert any("B2=1200" in section.content for section in result.sections)
    sales_row = next(
        section
        for section in result.sections
        if section.meta.get("sheet") == "销售数据" and section.meta.get("row") == 2
    )
    assert sales_row.meta["row"] == 2
    assert sales_row.meta["cell_addresses"] == ["A2", "B2"]
    assert sales_row.section_path == ["sales.xlsx", "销售数据", "第 2 行"]


def test_registry_resolves_and_parses_real_pptx_with_notes_and_tables() -> None:
    registry = ParserRegistry()
    parser = registry.resolve(PPTX_MIME)
    assert parser.parser_id == "pptx"

    result = registry.parse(_pptx_bytes(), filename="quarterly.pptx", mime=PPTX_MIME)

    assert result.metadata["format"] == "pptx"
    assert result.metadata["slides"] == 2
    assert result.metadata["populated_slides"] == 2
    assert any("演讲者备注" in section.content for section in result.sections)
    table_section = next(
        section for section in result.sections if section.meta["table_shape_count"]
    )
    assert "R2C2=42" in table_section.content
    assert table_section.meta["slide"] == 2


@pytest.mark.parametrize(
    ("mime", "filename"),
    [(XLSX_MIME, "broken.xlsx"), (PPTX_MIME, "broken.pptx")],
)
def test_corrupt_ooxml_is_stable_and_does_not_leak_library_errors(
    mime: str, filename: str
) -> None:
    with pytest.raises(ParserError) as caught:
        ParserRegistry().parse(b"not-an-ooxml-document", filename=filename, mime=mime)

    error = caught.value
    assert error.code == PARSER_CORRUPT
    assert error.to_sanitized_error()["detail"]["reason"] in {
        "xlsx-read-error",
        "pptx-read-error",
    }
    assert "BadZipFile" not in error.message
    assert "not-an-ooxml-document" not in error.message


@pytest.mark.parametrize(
    "mime", [XLSX_MIME, PPTX_MIME]
)
def test_empty_ooxml_is_file_empty(mime: str) -> None:
    with pytest.raises(ParserError) as caught:
        ParserRegistry().parse(b"", filename="empty.office", mime=mime)
    assert caught.value.code == FILE_EMPTY


@pytest.mark.parametrize(
    "mime",
    ["application/msword", "application/vnd.ms-excel", "application/vnd.ms-powerpoint"],
)
def test_legacy_office_mimes_remain_fail_closed(mime: str) -> None:
    with pytest.raises(ParserError) as caught:
        ParserRegistry().parse(b"legacy-bytes", filename="legacy.office", mime=mime)
    assert caught.value.code == MIME_UNSUPPORTED
